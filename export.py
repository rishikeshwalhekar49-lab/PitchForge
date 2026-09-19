"""Export an edited deck to .pptx (fully editable in PowerPoint / Google Slides / Keynote) or Markdown."""
import io

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


def to_pptx(deck) -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for s in deck["slides"]:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        tag = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12), Inches(0.5)).text_frame
        tag.text = f"{s['n']:02d} · {s['title'].upper()}"
        tag.paragraphs[0].runs[0].font.size = Pt(14)
        tag.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x4F, 0x46, 0xE5)
        head = slide.shapes.add_textbox(Inches(0.6), Inches(0.9), Inches(12), Inches(1.4)).text_frame
        head.word_wrap = True
        head.text = s["headline"]
        head.paragraphs[0].runs[0].font.size = Pt(34)
        head.paragraphs[0].runs[0].font.bold = True
        body = slide.shapes.add_textbox(Inches(0.6), Inches(2.5), Inches(12), Inches(4.5)).text_frame
        body.word_wrap = True
        for i, b in enumerate(s["bullets"]):
            p = body.paragraphs[0] if i == 0 else body.add_paragraph()
            p.text = f"•  {b}"
            p.font.size = Pt(20)
            p.space_after = Pt(12)
        slide.notes_slide.notes_text_frame.text = s.get("speaker_notes", "")
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def to_markdown(deck) -> str:
    out = [f"# {deck['input'].get('company') or 'Pitch'} — {deck['input']['idea']}\n"]
    for s in deck["slides"]:
        out.append(f"## {s['n']}. {s['title']}\n**{s['headline']}**\n")
        out += [f"- {b}" for b in s["bullets"]]
        out.append(f"\n> Notes: {s.get('speaker_notes', '')}\n")
    return "\n".join(out)
